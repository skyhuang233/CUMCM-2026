"""结果写出与表 1 / 表 2 / 表 3 摘要。

附件 5 官方模板缺失，按题面重建布局；拿到模板后只改本文件。
"""

from __future__ import annotations

import os
from typing import Sequence

import numpy as np
import openpyxl

from .data import INTERVAL_LABELS, INTERVAL_SEGS, SEG_LABELS
from .optimizer import LPResult
from .settlement import (  # noqa: F401
    cost_q2,
    cost_q3,
    emergency_intervals,
    merge_emergency_intervals,
)

# 表 1 指定的 10 分钟段（10:00-10:10 … 20:00-20:10）
TABLE1_SEGS = [60, 72, 84, 96, 108, 120]
TABLE1_LABELS = ["10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10"]


def interval_sums(x: np.ndarray, n_per: int = INTERVAL_SEGS) -> np.ndarray:
    """把逐段序列按 4 小时区间求和。"""
    x = np.asarray(x, dtype=float)
    return x.reshape(-1, n_per).sum(axis=1)


def summarize_q1(res: LPResult, price: np.ndarray, soc_init: float) -> dict:
    """表 1（购电）与表 2（充放电与储电量）所需的全部数值。"""
    price = np.asarray(price, dtype=float)
    return {
        "seg_labels": TABLE1_LABELS,
        "seg_purchase": res.G[TABLE1_SEGS],
        "daily_purchase": float(res.G.sum()),
        "purchase_cost": float(price @ res.G),
        "interval_labels": INTERVAL_LABELS,
        "charge": interval_sums(res.C),
        "discharge": interval_sums(res.D),
        "soc_start": float(soc_init),
        "soc_end": float(res.S[-1]),
    }


def write_result1(res: LPResult, soc_init: float, path: str = "results/result1.xlsx", label: str = "附件1") -> str:
    """写出 result1.xlsx：工作表『计划购电量』与『充放电量』。"""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["日期"] + SEG_LABELS)
    ws.append([label] + [round(float(v), 4) for v in res.G])

    ws2 = wb.create_sheet("充放电量")
    ws2.append(["时间段", "充电量", "放电量"])
    charge, discharge = interval_sums(res.C), interval_sums(res.D)
    for name, c, d in zip(INTERVAL_LABELS, charge, discharge):
        ws2.append([name, round(float(c), 4), round(float(d), 4)])
    ws2.append(["0:00 储电量", round(float(soc_init), 4), None])
    ws2.append(["24:00 储电量", round(float(res.S[-1]), 4), None])

    wb.save(path)
    return path


def print_summary(summary: dict) -> None:
    """终端打印表 1 / 表 2。"""
    print("表 1 计划购电量（kWh）")
    for name, val in zip(summary["seg_labels"], summary["seg_purchase"]):
        print(f"  {name:>12}  {val:10.2f}")
    print(f"  {'全天购电量':>12}  {summary['daily_purchase']:10.2f}")
    print(f"  {'全天购电费':>12}  {summary['purchase_cost']:10.2f} 元")
    print()
    print("表 2 充放电量与储电量（kWh）")
    print(f"  {'时间段':>12}  {'充电量':>10}  {'放电量':>10}")
    for name, c, d in zip(summary["interval_labels"], summary["charge"], summary["discharge"]):
        print(f"  {name:>12}  {c:10.2f}  {d:10.2f}")
    print(f"  {'0:00 储电量':>12}  {summary['soc_start']:10.2f}")
    print(f"  {'24:00 储电量':>12}  {summary['soc_end']:10.2f}")


# ---------------------------------------------------------------------------
# 问 2：逐日结果写出与表 1 / 表 2 / 表 3 摘要
# ---------------------------------------------------------------------------

CHARGE_COLUMNS = [f"{name} {kind}" for name in INTERVAL_LABELS for kind in ("充电量", "放电量")]


def summarize_day(day, price: np.ndarray) -> dict:
    """一天的表 1（购电）、表 2（执行充放电与储电量）、表 3（紧急购电区间）数值。

    `day` 需具备 G0, C, D, S, E, soc_start 属性（`run_q2.DayResult` 或等价对象）。
    """
    price = np.asarray(price, dtype=float)
    plan_cost, emergency_cost = cost_q2(price, day.G0, day.E)
    point = getattr(day, "point_solution", None)
    out = {
        "date": day.day,
        "seg_labels": TABLE1_LABELS,
        "seg_purchase": np.asarray(day.G0)[TABLE1_SEGS],
        "daily_purchase": float(np.sum(day.G0)),
        "plan_cost": plan_cost,
        "emergency_cost": emergency_cost,
        "purchase_cost": plan_cost + emergency_cost,
        "emergency_kwh": float(np.sum(day.E)),
        "interval_labels": INTERVAL_LABELS,
        "charge": interval_sums(day.C),
        "discharge": interval_sums(day.D),
        "soc_start": float(day.soc_start),
        "soc_end": float(np.asarray(day.S)[-1]),
        "emergency": emergency_intervals(day.day, day.E),
    }
    if point is not None:
        n = len(day.G0)
        out["point_charge"] = interval_sums(point.C[:n])
        out["point_discharge"] = interval_sums(point.D[:n])
        out["point_soc_end"] = float(point.S[n - 1])
    return out


def print_day_summary(summary: dict) -> None:
    """终端打印某一天的表 1 / 表 2 / 表 3。"""
    d = summary["date"]
    print(f"=== {d.isoformat()} ===")
    print("表 1 计划购电量（kWh）")
    for name, val in zip(summary["seg_labels"], summary["seg_purchase"]):
        print(f"  {name:>12}  {val:10.2f}")
    print(f"  {'全天购电量':>12}  {summary['daily_purchase']:10.2f}")
    print(
        f"  {'全天购电费':>12}  {summary['purchase_cost']:10.2f} 元"
        f"（计划 {summary['plan_cost']:.2f} + 紧急 {summary['emergency_cost']:.2f}）"
    )
    print("表 2 充放电量与储电量（kWh，执行值）")
    print(f"  {'时间段':>12}  {'充电量':>10}  {'放电量':>10}")
    for name, c, dis in zip(summary["interval_labels"], summary["charge"], summary["discharge"]):
        print(f"  {name:>12}  {c:10.2f}  {dis:10.2f}")
    print(f"  {'0:00 储电量':>12}  {summary['soc_start']:10.2f}")
    print(f"  {'24:00 储电量':>12}  {summary['soc_end']:10.2f}")
    print("表 3 紧急购电量（kWh）")
    if not summary["emergency"]:
        print("  （当日无紧急购电）")
    for label, kwh in summary["emergency"]:
        print(f"  {label:>12}  {kwh:10.2f}")
    if "point_charge" in summary:
        print("（对照）点预测解的充放电量与 24:00 储电量")
        for name, c, dis in zip(
            summary["interval_labels"], summary["point_charge"], summary["point_discharge"]
        ):
            print(f"  {name:>12}  {c:10.2f}  {dis:10.2f}")
        print(f"  {'24:00 储电量':>12}  {summary['point_soc_end']:10.2f}")
    print()


def write_result2(days: Sequence, path: str = "results/result2.xlsx") -> str:
    """写出 result2.xlsx：『计划购电量』『充放电量』『紧急购电量』三张工作表。"""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["日期"] + SEG_LABELS)
    for day in days:
        ws.append([day.day.isoformat()] + [round(float(v), 4) for v in day.G0])

    ws2 = wb.create_sheet("充放电量")
    ws2.append(["日期"] + CHARGE_COLUMNS + ["0:00 储电量", "24:00 储电量"])
    for day in days:
        charge, discharge = interval_sums(day.C), interval_sums(day.D)
        row: list = [day.day.isoformat()]
        for c, dis in zip(charge, discharge):
            row += [round(float(c), 4), round(float(dis), 4)]
        row += [round(float(day.soc_start), 4), round(float(np.asarray(day.S)[-1]), 4)]
        ws2.append(row)

    ws3 = wb.create_sheet("紧急购电量")
    ws3.append(["日期", "紧急购电时间段", "紧急购电量"])
    for day in days:
        first = True
        for label, kwh in emergency_intervals(day.day, day.E):
            ws3.append([day.day.isoformat() if first else None, label, round(float(kwh), 4)])
            first = False

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# 问 3：调整购电量的写出、四日摘要与变体对比表
# ---------------------------------------------------------------------------


def summarize_day_q3(day, price: np.ndarray) -> dict:
    """一天的表 1（计划与调整购电）、表 2（执行充放电与储电量）、表 3（紧急购电区间）。

    `day` 需具备 G0, Ga, C, D, S, E, soc_start 属性（`run_q3.DayResult` 或等价对象）。
    表 1 的购电量与购电费按**提交的** $G^a$ 与问 3 的四项分解计费。
    """
    price = np.asarray(price, dtype=float)
    G0 = np.asarray(day.G0, dtype=float)
    Ga = np.asarray(day.Ga, dtype=float)
    cost = cost_q3(price, G0, Ga, day.E)
    return {
        "date": day.day,
        "seg_labels": TABLE1_LABELS,
        "seg_plan": G0[TABLE1_SEGS],
        "seg_purchase": Ga[TABLE1_SEGS],
        "daily_plan": float(G0.sum()),
        "daily_purchase": float(Ga.sum()),
        "cost": cost,
        "purchase_cost": cost["total"],
        "emergency_kwh": float(np.sum(day.E)),
        "adjust_kwh": float(np.abs(Ga - G0).sum()),
        "interval_labels": INTERVAL_LABELS,
        "charge": interval_sums(day.C),
        "discharge": interval_sums(day.D),
        "soc_start": float(day.soc_start),
        "soc_end": float(np.asarray(day.S)[-1]),
        "emergency": emergency_intervals(day.day, day.E),
    }


def print_day_summary_q3(summary: dict) -> None:
    """终端打印问 3 某一天的表 1 / 表 2 / 表 3。"""
    c = summary["cost"]
    print(f"=== {summary['date'].isoformat()} ===")
    print("表 1 购电量（kWh）")
    print(f"  {'时间段':>12}  {'计划购电量':>12}  {'调整购电量':>12}")
    for name, g0, ga in zip(summary["seg_labels"], summary["seg_plan"], summary["seg_purchase"]):
        print(f"  {name:>12}  {g0:12.2f}  {ga:12.2f}")
    print(f"  {'全天购电量':>12}  {summary['daily_plan']:12.2f}  {summary['daily_purchase']:12.2f}")
    print(
        f"  {'全天购电费':>12}  {c['total']:12.2f} 元"
        f"（计划 {c['plan']:.2f} + 减购违约 {c['curtail_penalty']:.2f}"
        f" + 增购 {c['extra']:.2f} + 紧急 {c['emergency']:.2f}）"
    )
    print("表 2 充放电量与储电量（kWh，执行值）")
    print(f"  {'时间段':>12}  {'充电量':>10}  {'放电量':>10}")
    for name, ch, dis in zip(summary["interval_labels"], summary["charge"], summary["discharge"]):
        print(f"  {name:>12}  {ch:10.2f}  {dis:10.2f}")
    print(f"  {'0:00 储电量':>12}  {summary['soc_start']:10.2f}")
    print(f"  {'24:00 储电量':>12}  {summary['soc_end']:10.2f}")
    print("表 3 紧急购电量（kWh）")
    if not summary["emergency"]:
        print("  （当日无紧急购电）")
    for label, kwh in summary["emergency"]:
        print(f"  {label:>12}  {kwh:10.2f}")
    print()


def write_result3(days: Sequence, path: str = "results/result3.xlsx") -> str:
    """写出 result3.xlsx：『计划购电量』『调整购电量』『充放电量』『紧急购电量』四张工作表。"""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["日期"] + SEG_LABELS)
    for day in days:
        ws.append([day.day.isoformat()] + [round(float(v), 4) for v in day.G0])

    ws2 = wb.create_sheet("调整购电量")
    ws2.append(["日期"] + SEG_LABELS)
    for day in days:
        ws2.append([day.day.isoformat()] + [round(float(v), 4) for v in day.Ga])

    ws3 = wb.create_sheet("充放电量")
    ws3.append(["日期"] + CHARGE_COLUMNS + ["0:00 储电量", "24:00 储电量"])
    for day in days:
        charge, discharge = interval_sums(day.C), interval_sums(day.D)
        row: list = [day.day.isoformat()]
        for ch, dis in zip(charge, discharge):
            row += [round(float(ch), 4), round(float(dis), 4)]
        row += [round(float(day.soc_start), 4), round(float(np.asarray(day.S)[-1]), 4)]
        ws3.append(row)

    ws4 = wb.create_sheet("紧急购电量")
    ws4.append(["日期", "紧急购电时间段", "紧急购电量"])
    for day in days:
        first = True
        for label, kwh in emergency_intervals(day.day, day.E):
            ws4.append([day.day.isoformat() if first else None, label, round(float(kwh), 4)])
            first = False

    wb.save(path)
    return path


def variant_table(rows: Sequence[dict]) -> str:
    """变体对比表：每行一个 `--issues` 子集。"""
    head = (
        f"{'预报时刻':<16} {'全年总费用':>15} {'计划费':>15} {'减购违约费':>13}"
        f" {'增购费':>13} {'紧急费':>13} {'紧急购电kWh':>13} {'平均调整量':>11}"
    )
    lines = [head]
    for r in rows:
        lines.append(
            f"{r['label']:<16} {r['total_cost']:>15.2f} {r['plan_cost']:>15.2f}"
            f" {r['curtail_cost']:>13.2f} {r['extra_cost']:>13.2f}"
            f" {r['emergency_cost']:>13.2f} {r['emergency_kwh']:>13.2f}"
            f" {r['mean_adjust_kwh']:>11.2f}"
        )
    return "\n".join(lines)
