import numpy as np
import openpyxl
import pytest
from datetime import date

from src.data import SOC_INIT, SOC_MAX, SOC_MIN, T
from src.results import write_result2
from src.run_q2 import Params, load_bundle, run_period, validate


@pytest.fixture(scope="module")
def short_run():
    bundle = load_bundle()
    res = run_period(
        date(2025, 1, 1),
        date(2025, 2, 2),
        Params(step_minutes=60),
        bundle,
        record_from=date(2025, 2, 1),
        soc_init=SOC_INIT,
    )
    return bundle, res


def test_records_only_from_record_from(short_run):
    _, res = short_run
    assert [d.day for d in res.days] == [
        date(2025, 2, 1),
        date(2025, 2, 2),
    ]


def test_soc_is_continuous_across_days_and_in_bounds(short_run):
    _, res = short_run
    for prev, cur in zip(res.days, res.days[1:]):
        assert abs(cur.soc_start - prev.soc_end) < 1e-9
    for day in res.days:
        assert day.S.min() >= SOC_MIN - 1e-6 and day.S.max() <= SOC_MAX + 1e-6
        assert np.minimum(day.C, day.D).max() < 1e-6
        assert np.minimum(day.E, day.W).max() < 1e-9


def test_cost_decomposition_matches_total(short_run):
    bundle, res = short_run
    price = bundle.att1.price
    for day in res.days:
        assert abs(day.plan_cost - float(price @ day.G0)) < 1e-6
        assert abs(day.emergency_cost - 5.0 * float(price @ day.E)) < 1e-6
        assert abs(day.total_cost - (day.plan_cost + day.emergency_cost)) < 1e-9
    assert abs(res.total_cost - sum(d.total_cost for d in res.days)) < 1e-6


def test_validate_passes_on_a_real_run(short_run):
    bundle, res = short_run
    validate(res, bundle)  # 平衡、E·W=0、SOC 界与跨日连续、费用分解


def test_validate_rejects_a_broken_soc_chain(short_run):
    import copy

    bundle, res = short_run
    broken = copy.deepcopy(res)
    broken.days[1].soc_start += 100.0
    with pytest.raises(AssertionError):
        validate(broken, bundle)


def test_48h_increment_runs_as_a_distinct_supported_configuration():
    bundle = load_bundle()
    res = run_period(
        date(2025, 1, 1), date(2025, 2, 1),
        Params(horizon_days=2, m_scen=2), bundle,
        record_from=date(2025, 2, 1), soc_init=SOC_INIT,
    )
    assert len(res.days) == 1
    validate(res, bundle)


def test_q4_2_freezes_purchase_but_refreshes_value_four_times(monkeypatch):
    from src.data import load_attachment4
    from src.executor import DPValueExecutor
    from src.forecast_price import PriceForecaster

    bundle = load_bundle()
    dates, prices = load_attachment4()
    source = PriceForecaster(dates, prices, bundle.att1, branch="q4_2")
    prepares = []
    original = DPValueExecutor.prepare

    def spy(self, t0=0, hbar=None):
        prepares.append(t0)
        return original(self, t0, hbar)

    monkeypatch.setattr(DPValueExecutor, "prepare", spy)
    res = run_period(
        date(2025, 1, 1), date(2025, 2, 1), Params(m_scen=2), bundle,
        record_from=date(2025, 2, 1), price_source=source,
    )
    assert len(res.days) == 1
    assert prepares[-4:] == [0, 36, 72, 108]
    # The only submitted quantity is G0; Q4-2 has no adjustment record/API.
    assert res.days[0].G0.shape == (T,)


def test_48h_q4_2_intraday_values_only_cover_remaining_today(monkeypatch):
    import src.run_q2 as q2
    from src.data import load_attachment4
    from src.forecast_price import PriceForecaster

    bundle = load_bundle()
    dates, prices = load_attachment4()
    source = PriceForecaster(dates, prices, bundle.att1, branch="q4_2")
    lengths = []
    original = q2.future_cost

    def spy(price, load, pv, g, *args, **kwargs):
        lengths.append((len(g), np.asarray(load).shape[1], np.asarray(price).shape[-1]))
        return original(price, load, pv, g, *args, **kwargs)

    monkeypatch.setattr(q2, "future_cost", spy)
    run_period(
        date(2025, 1, 1), date(2025, 2, 1),
        Params(horizon_days=2, m_scen=2), bundle,
        record_from=date(2025, 2, 1), price_source=source,
    )
    # 0:00 evaluation plus 6/12/18 each use only the remaining current day;
    # V(e) alone prices the following whole day.
    assert lengths[-4:] == [(T, T, T), (108, 108, 108), (72, 72, 72), (36, 36, 36)]


def test_default_q2_builds_one_hbar_and_skips_unrequested_point_lp(monkeypatch):
    import src.run_q2 as q2

    bundle = load_bundle()
    point_calls = []
    hbar_calls = []
    original_future = q2.future_cost
    original_point = q2.solve_saa_point

    def point_spy(*args, **kwargs):
        point_calls.append(1)
        return original_point(*args, **kwargs)

    def future_spy(*args, **kwargs):
        hbar_calls.append(1)
        return original_future(*args, **kwargs)

    monkeypatch.setattr(q2, "solve_saa_point", point_spy)
    monkeypatch.setattr(q2, "future_cost", future_spy)
    run_period(
        date(2025, 1, 1), date(2025, 2, 1), Params(m_scen=2), bundle,
        record_from=date(2025, 2, 1), point_days=set(),
    )
    assert point_calls == []
    assert len(hbar_calls) == 1


def test_checkpoint_resume_matches_an_uninterrupted_q2_run(tmp_path, monkeypatch):
    """A stop after a completed scoring day must be exactly resumable."""
    import src.run_q2 as q2

    bundle = load_bundle()
    params = Params(m_scen=1)
    checkpoint = tmp_path / "q2.pkl"
    original = q2.future_cost
    calls = 0

    def interrupted(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:  # Feb 1 was checkpointed; interrupt on Feb 2.
            raise RuntimeError("intentional interruption")
        return original(*args, **kwargs)

    monkeypatch.setattr(q2, "future_cost", interrupted)
    with pytest.raises(RuntimeError, match="intentional"):
        run_period(
            date(2025, 1, 1), date(2025, 2, 2), params, bundle,
            record_from=date(2025, 2, 1), checkpoint_path=checkpoint,
        )
    assert checkpoint.exists()
    monkeypatch.setattr(q2, "future_cost", original)

    resumed = run_period(
        date(2025, 1, 1), date(2025, 2, 2), params, bundle,
        record_from=date(2025, 2, 1), checkpoint_path=checkpoint, resume=True,
    )
    clean = run_period(
        date(2025, 1, 1), date(2025, 2, 2), params, bundle,
        record_from=date(2025, 2, 1),
    )
    assert resumed.total_cost == clean.total_cost
    for got, expected in zip(resumed.days, clean.days):
        assert got.day == expected.day
        assert np.array_equal(got.G0, expected.G0)
        assert np.array_equal(got.S, expected.S)


def test_main_never_carries_calibration_dispatch_into_formal_run(monkeypatch):
    import src.run_q2 as q2

    captured = []

    def fake_run(*args, **kwargs):
        captured.append(args[2])
        return q2.PeriodResult()

    monkeypatch.setattr(q2, "run_period", fake_run)
    monkeypatch.setattr(q2, "print_period_summary", lambda *args, **kwargs: None)
    q2.main(["--tune", "--skip-tuning", "--end", "2025-02-01", "--no-write"])
    assert captured and captured[0].calibration is False


def test_write_result2_layout(short_run, tmp_path):
    _, res = short_run
    path = write_result2(res.days, str(tmp_path / "result2.xlsx"))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["计划购电量", "充放电量", "紧急购电量"]
    plan = wb["计划购电量"]
    assert plan.max_row == len(res.days) + 1
    assert plan.max_column == T + 1
    assert plan.cell(1, 2).value == "0:10" and plan.cell(1, T + 1).value == "0:00+1"
    assert plan.cell(2, 1).value == res.days[0].day.isoformat()

    charge = wb["充放电量"]
    assert charge.max_column == 2 * 6 + 2 + 1
    assert charge.cell(1, 2).value == "0:00-4:00 充电量"
    assert charge.cell(1, charge.max_column).value == "24:00 储电量"

    emergency = wb["紧急购电量"]
    assert [c.value for c in emergency[1]] == ["日期", "紧急购电时间段", "紧急购电量"]
    # 同一日期多区间时日期只在首行填写
    seen_dates = [emergency.cell(r, 1).value for r in range(2, emergency.max_row + 1)]
    assert all(v is None or isinstance(v, str) for v in seen_dates)
