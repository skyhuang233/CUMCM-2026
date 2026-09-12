import numpy as np
import openpyxl
import pytest
from datetime import date

from src.data import SOC_INIT, SOC_MAX, SOC_MIN, T, load_attachment4
from src.forecast_price import PerfectPriceSource, PriceForecaster
from src.executor import DPValueExecutor
from src.results import write_result2, write_result3
from src.run_q2 import Params as Q2Params
from src.run_q2 import load_bundle as load_bundle_q2
from src.run_q2 import run_period as run_period_q2
from src.run_q2 import validate as validate_q2
from src.run_q3 import Params as Q3Params
from src.run_q3 import load_bundle as load_bundle_q3
from src.run_q3 import run_period as run_period_q3
from src.run_q3 import validate as validate_q3

D0, D1 = date(2025, 1, 1), date(2025, 2, 2)
PARAMS2 = Q2Params(k_load=6, k_pv=7, step_minutes=60)
PARAMS3 = Q3Params(step_minutes=60)


@pytest.fixture(scope="module")
def att4():
    return load_attachment4()


@pytest.fixture(scope="module")
def bundle2():
    return load_bundle_q2()


@pytest.fixture(scope="module")
def bundle3():
    return load_bundle_q3()


def price_source(att4, att1, k=4, branch="q4_2"):
    dates, prices = att4
    return PriceForecaster(dates, prices, att1, k=k, branch=branch)


@pytest.fixture(scope="module")
def run2(bundle2, att4):
    res = run_period_q2(
        D0,
        D1,
        PARAMS2,
        bundle2,
        record_from=date(2025, 2, 1),
        soc_init=SOC_INIT,
        price_source=price_source(att4, bundle2.att1),
    )
    return bundle2, res


@pytest.fixture(scope="module")
def run3(bundle3, att4):
    res = run_period_q3(
        D0,
        D1,
        PARAMS3,
        bundle3,
        record_from=date(2025, 2, 1),
        soc_init=SOC_INIT,
        price_source=price_source(att4, bundle3.att1, branch="q4_3"),
    )
    return bundle3, res


def test_q4_2_settles_at_attachment4_truth(run2, att4):
    """每日费用必须按附件 4 当日真值结算，而不是附件 1 的常数电价。"""
    bundle, res = run2
    dates, prices = att4
    for day in res.days:
        truth = prices[dates.index(day.day)]
        assert day.price is not None and np.allclose(day.price, truth)
        assert abs(day.plan_cost - float(truth @ day.G0)) < 1e-6
        assert abs(day.emergency_cost - 5.0 * float(truth @ day.E)) < 1e-6
        assert not np.allclose(truth, bundle.att1.price)
    validate_q2(res, bundle)


def test_q4_3_settles_at_attachment4_truth(run3, att4):
    bundle, res = run3
    dates, prices = att4
    for day in res.days:
        truth = prices[dates.index(day.day)]
        assert np.allclose(day.price, truth)
        assert abs(day.emergency_cost - 5.0 * float(truth @ day.E)) < 1e-6
        assert abs(day.plan_only_cost - float(truth @ day.G0)) < 1e-6
    validate_q3(res, bundle, PARAMS3.issues)


def test_q4_runs_keep_the_physical_invariants(run2, run3):
    for _, res in (run2, run3):
        for day in res.days:
            assert day.S.min() >= SOC_MIN - 1e-6 and day.S.max() <= SOC_MAX + 1e-6
            assert np.minimum(day.C, day.D).max() < 1e-6
            assert np.minimum(day.E, day.W).max() < 1e-9
            assert day.G0.min() >= -1e-9


def test_price_paths_are_available_to_the_q4_run(run2):
    """二月运行必须真正产生非空计划，而不是用空的一月循环掩盖集成错误。"""
    _, result = run2
    assert len(result.days) == 2
    assert all(np.isfinite(day.G0).all() and day.G0.sum() > 0 for day in result.days)


def test_no_future_price_reaches_any_decision(bundle2, bundle3, att4):
    """把末日之后的附件 4 电价毒化成 NaN：$G^0$、$G^a$ 与费用必须逐位不变。"""
    dates, prices = att4
    i = dates.index(D1)
    poisoned = prices.copy()
    poisoned[i + 1 :] = np.nan

    for bundle, params, runner, branch in (
        (bundle2, PARAMS2, run_period_q2, "q4_2"),
        (bundle3, PARAMS3, run_period_q3, "q4_3"),
    ):
        clean = runner(
            D0, D1, params, bundle, soc_init=SOC_INIT,
            price_source=PriceForecaster(dates, prices, bundle.att1, k=4, branch=branch),
        )
        guarded = runner(
            D0, D1, params, bundle, soc_init=SOC_INIT,
            price_source=PriceForecaster(dates, poisoned, bundle.att1, k=4, branch=branch),
        )
        for a, b in zip(clean.days, guarded.days):
            assert np.array_equal(a.G0, b.G0)
            assert np.array_equal(getattr(a, "Ga", a.G0), getattr(b, "Ga", b.G0))
        assert clean.total_cost == guarded.total_cost


def test_no_decision_sees_a_price_segment_beyond_the_segment_it_decides(
    bundle3, att4, monkeypatch
):
    """记录每次电价预测与每次 DP 执行的段号：段 t 的决策之前，当天用过的
    已观测段数不得超过 t+1（即从不读取段 $\\ge t+1$ 的电价）。"""
    dates, prices = att4
    events: list[tuple[str, date | None, int]] = []
    predict = PriceForecaster.predict
    step = DPValueExecutor.step

    def spy_predict(self, d, t_now=0, observed=None):
        events.append(("predict", d, int(t_now)))
        return predict(self, d, t_now, observed)

    def spy_step(self, t, *a, **kw):
        events.append(("step", None, int(t)))
        return step(self, t, *a, **kw)

    monkeypatch.setattr(PriceForecaster, "predict", spy_predict)
    monkeypatch.setattr(DPValueExecutor, "step", spy_step)
    run_period_q3(
        D0,
        date(2025, 2, 1),
        PARAMS3,
        bundle3,
        soc_init=SOC_INIT,
        price_source=PriceForecaster(dates, prices, bundle3.att1, k=4, branch="q4_3"),
    )
    assert any(kind == "predict" for kind, _, _ in events)
    assert any(kind == "step" for kind, _, _ in events)
    seen: dict[date, int] = {}
    current: date | None = None
    for kind, d, t in events:
        if kind == "predict":
            current = d
            seen[d] = max(seen.get(d, 0), t)
        else:
            # 段 t 的 DP 决策之前，当天最多只观测到段 t（即 t+1 个段）
            assert seen.get(current, 0) <= t + 1, f"段 {t} 之前读到了 {seen[current]} 段电价"


def test_perfect_price_variant_is_reported_only_as_an_information_reference(bundle2, att4):
    dates, prices = att4
    end = date(2025, 2, 2)
    forecast = run_period_q2(
        D0, end, PARAMS2, bundle2, record_from=date(2025, 2, 1), soc_init=SOC_INIT,
        price_source=PriceForecaster(dates, prices, bundle2.att1, k=4),
    )
    perfect = run_period_q2(
        D0, end, PARAMS2, bundle2, record_from=date(2025, 2, 1), soc_init=SOC_INIT,
        price_source=PerfectPriceSource(dates, prices),
    )
    # This still uses the approximate SAA/DP policy, so price information
    # alone is not an offline lower bound and need not dominate every run.
    assert forecast.days and perfect.days
    assert forecast.total_cost > 0 and perfect.total_cost > 0


def test_result4_layouts_match_result2_and_result3(run2, run3, tmp_path):
    _, res2 = run2
    _, res3 = run3
    p2 = write_result2(res2.days, str(tmp_path / "result4-2.xlsx"))
    p3 = write_result3(res3.days, str(tmp_path / "result4-3.xlsx"))
    wb2 = openpyxl.load_workbook(p2)
    assert wb2.sheetnames == ["计划购电量", "充放电量", "紧急购电量"]
    assert wb2["计划购电量"].max_row == len(res2.days) + 1
    assert wb2["计划购电量"].max_column == T + 1
    wb3 = openpyxl.load_workbook(p3)
    assert wb3.sheetnames == ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    assert wb3["调整购电量"].max_row == len(res3.days) + 1


def test_q4_cli_propagates_all_configs_to_both_and_constant_runs(monkeypatch, bundle2, bundle3, att4):
    import src.run_q4 as q4

    seen = []

    class Fake:
        days = []
        total_cost = 1.0
        def by_date(self, _): return None

    monkeypatch.setattr(q4, "load_attachment4", lambda: att4)
    monkeypatch.setattr(q4, "load_bundle_q2", lambda: bundle2)
    monkeypatch.setattr(q4, "load_bundle_q3", lambda: bundle3)
    monkeypatch.setattr(q4, "run_q4_2", lambda bundle, source, params, **kw: seen.append(("q4_2", params, source)) or Fake())
    monkeypatch.setattr(q4, "run_q4_3", lambda bundle, source, params, **kw: seen.append(("q4_3", params, source)) or Fake())
    monkeypatch.setattr(q4, "run_period_q2", lambda *a, **kw: seen.append(("const2", a[2], None)) or Fake())
    monkeypatch.setattr(q4, "run_period_q3", lambda *a, **kw: seen.append(("const3", a[2], None)) or Fake())
    monkeypatch.setattr(q4, "validate_q2", lambda *a: None)
    monkeypatch.setattr(q4, "validate_q3", lambda *a: None)
    monkeypatch.setattr(q4, "print_period_summary_q2", lambda *a, **kw: None)
    monkeypatch.setattr(q4, "print_period_summary_q3", lambda *a, **kw: None)
    q4.main(["--which", "both", "--with-constant", "--no-write", "--end", "2025-02-01",
             "--horizon-days", "2", "--candidate-reeval", "--legacy-means",
             "--same-type-scenarios", "--legacy-price"])
    assert [x[0] for x in seen] == ["q4_2", "q4_3", "const2", "const3"]
    for _, params, _ in seen:
        assert params.horizon_days == 2 and params.candidate_reeval
        assert not params.reference_baseline and params.scenario_same_type
        assert params.calibration is False
    assert seen[0][2].mode == "legacy_level"


def test_q4_cli_default_does_not_tune_or_use_legacy_price(monkeypatch, bundle2, att4):
    import src.run_q4 as q4
    called = []

    class Fake:
        days = []
        total_cost = 1.0
        def by_date(self, _): return None

    monkeypatch.setattr(q4, "load_attachment4", lambda: att4)
    monkeypatch.setattr(q4, "load_bundle_q2", lambda: bundle2)
    monkeypatch.setattr(q4, "run_q4_2", lambda bundle, source, params, **kw: called.append((source, params)) or Fake())
    monkeypatch.setattr(q4, "print_period_summary_q2", lambda *a, **kw: None)
    q4.main(["--which", "2", "--no-write", "--end", "2025-02-01"])
    assert called[0][0].mode == "baseline"
    assert called[0][1].calibration is False


def test_q4_tune_uses_legacy_price_but_formal_params_are_not_calibration(monkeypatch, bundle2, att4):
    import sys
    from types import SimpleNamespace
    import src.run_q4 as q4

    captured = []

    class Fake:
        days = []
        total_cost = 1.0
        def by_date(self, _): return None

    monkeypatch.setitem(sys.modules, "src.tuning", SimpleNamespace(
        select_kp=lambda *a, **kw: SimpleNamespace(k_price=9, window="January", cost=1.0)
    ))
    monkeypatch.setattr(q4, "load_attachment4", lambda: att4)
    monkeypatch.setattr(q4, "load_bundle_q2", lambda: bundle2)
    monkeypatch.setattr(q4, "run_q4_2", lambda bundle, source, params, **kw: captured.append((source, params)) or Fake())
    monkeypatch.setattr(q4, "print_period_summary_q2", lambda *a, **kw: None)
    q4.main(["--which", "2", "--tune", "--no-write", "--end", "2025-02-01"])
    assert captured[0][0].k == 9 and captured[0][0].mode == "legacy_level"
    assert captured[0][1].calibration is False
